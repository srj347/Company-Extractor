"""Incremental Excel output + a resume checkpoint.

ExcelWriter saves after every page (crash-safe) and de-duplicates by profile
URL. CheckpointStore records the last completed page so a re-run resumes.
"""

from __future__ import annotations

import json
from pathlib import Path

from openpyxl import Workbook, load_workbook

from .models import Company

HEADERS = [
    "Company Name",
    "Industry",
    "Headquarter Location",
    "Overview",
    "Logo URL",
    "Profile URL",
    "Source Page",
    "Scraped At",
]
PROFILE_URL_COL = HEADERS.index("Profile URL")  # 0-based index used for dedup seed


class ExcelWriter:
    def __init__(self, path: Path) -> None:
        self.path = Path(path)
        self.seen: set[str] = set()

        if self.path.exists():
            # If the existing workbook has an older header schema, archive it
            # and start fresh — avoids appending into wrong columns.
            existing = load_workbook(self.path)
            existing_headers = [c.value for c in existing.active[1]] if existing.active.max_row else []
            if existing_headers[: len(HEADERS)] != HEADERS:
                archived = self.path.with_name(
                    self.path.stem + ".legacy" + self.path.suffix
                )
                # If the archive name is taken, append a counter.
                i = 1
                while archived.exists():
                    archived = self.path.with_name(
                        f"{self.path.stem}.legacy.{i}{self.path.suffix}"
                    )
                    i += 1
                self.path.rename(archived)
                print(
                    f"[storage] Schema changed — archived old workbook to {archived.name}.",
                    flush=True,
                )
                self.wb = Workbook()
                self.ws = self.wb.active
                self.ws.title = "Companies"
                self.ws.append(HEADERS)
                self._save()
            else:
                self.wb = existing
                self.ws = self.wb.active
                # Seed de-dup set from the existing Profile URL column.
                for row in self.ws.iter_rows(min_row=2, values_only=True):
                    if row and len(row) > PROFILE_URL_COL and row[PROFILE_URL_COL]:
                        self.seen.add(str(row[PROFILE_URL_COL]))
        else:
            self.wb = Workbook()
            self.ws = self.wb.active
            self.ws.title = "Companies"
            self.ws.append(HEADERS)
            self._save()

    def append(self, companies: list[Company]) -> int:
        added = 0
        for c in companies:
            key = c.profile_url or f"{c.name}|{c.hq_location}"
            if key in self.seen:
                continue
            self.seen.add(key)
            self.ws.append(c.as_row())
            added += 1
        self._save()
        return added

    @property
    def total_rows(self) -> int:
        return max(0, self.ws.max_row - 1)

    def _save(self) -> None:
        self.wb.save(self.path)


class CheckpointStore:
    def __init__(self, path: Path) -> None:
        self.path = Path(path)

    def last_completed_page(self) -> int:
        if self.path.exists():
            try:
                return int(json.loads(self.path.read_text()).get("last_completed_page", 0))
            except Exception:
                return 0
        return 0

    def save(self, page: int) -> None:
        self.path.write_text(json.dumps({"last_completed_page": page}))
