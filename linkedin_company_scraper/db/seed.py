"""Seed the PostgreSQL database from the scraped Excel file.

Usage:
    python -m linkedin_company_scraper.db.seed                          # dry-run (default)
    python -m linkedin_company_scraper.db.seed --execute                # actually insert
    python -m linkedin_company_scraper.db.seed --file other.xlsx        # custom file
    python -m linkedin_company_scraper.db.seed --create-tables          # run schema.sql first
    python -m linkedin_company_scraper.db.seed --normalize-only         # just run LLM normalization, save cache, exit

Env vars (or .env):
    DATABASE_URL   postgresql://user:pass@host:5432/dbname
    OPENAI_API_KEY sk-...
"""

from __future__ import annotations

import argparse
import sys
from datetime import datetime, timezone
from pathlib import Path

import psycopg2
from dotenv import load_dotenv
from openpyxl import load_workbook

from .normalizer import Normalizer

PROJECT_ROOT = Path(__file__).resolve().parent.parent.parent
SCHEMA_PATH = Path(__file__).resolve().parent / "schema.sql"

load_dotenv(PROJECT_ROOT / ".env")


# ── Excel parsing ─────────────────────────────────────────────────

def read_excel(path: Path) -> list[dict]:
    """Read companies.xlsx into a list of dicts keyed by header name."""
    wb = load_workbook(path, read_only=True)
    ws = wb.active
    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        return []
    headers = [str(h).strip() for h in rows[0]]
    records = []
    for row in rows[1:]:
        rec = {h: (v if v is not None else None) for h, v in zip(headers, row)}
        records.append(rec)
    wb.close()
    return records


# ── DB helpers ────────────────────────────────────────────────────

def create_tables(conn) -> None:
    sql = SCHEMA_PATH.read_text(encoding="utf-8")
    with conn.cursor() as cur:
        cur.execute(sql)
    conn.commit()
    print("[db] Tables created (or already exist).", flush=True)


def get_or_create_industry(cur, name: str | None) -> int | None:
    if not name:
        return None
    cur.execute("SELECT id FROM industry WHERE name = %s", (name,))
    row = cur.fetchone()
    if row:
        return row[0]
    cur.execute("INSERT INTO industry (name) VALUES (%s) RETURNING id", (name,))
    return cur.fetchone()[0]


def get_or_create_location(cur, city: str | None, state: str | None,
                           country: str | None) -> int | None:
    if not city and not state and not country:
        return None
    cur.execute(
        """SELECT id FROM location
           WHERE city IS NOT DISTINCT FROM %s
             AND state IS NOT DISTINCT FROM %s
             AND country IS NOT DISTINCT FROM %s""",
        (city, state, country),
    )
    row = cur.fetchone()
    if row:
        return row[0]
    cur.execute(
        """INSERT INTO location (city, state, country) VALUES (%s, %s, %s)
           RETURNING id""",
        (city, state, country),
    )
    return cur.fetchone()[0]


def company_exists(cur, source: str, platform: str) -> bool:
    cur.execute(
        "SELECT 1 FROM company WHERE source = %s AND platform = %s",
        (source, platform),
    )
    return cur.fetchone() is not None


def insert_company(cur, *, name, overview, industry_id, hq_location_id,
                   logo, website, platform, source, now) -> int:
    cur.execute(
        """INSERT INTO company
           (name, overview, industry_id, hq_location_id, logo, website,
            platform, source, created_at, updated_at)
           VALUES (%s, %s, %s, %s, %s, %s, %s, %s, %s, %s)
           RETURNING id""",
        (name, overview, industry_id, hq_location_id, logo, website,
         platform, source, now, now),
    )
    return cur.fetchone()[0]


# ── Main seeding logic ────────────────────────────────────────────

def seed(records: list[dict], conn, *, dry_run: bool = True) -> dict:
    """Normalize, de-dup, and insert records. Returns summary stats."""

    # 1. Collect unique raw industries and locations for batch LLM normalization.
    raw_industries = sorted({r["Industry"] for r in records if r.get("Industry")})
    raw_locations = sorted({r["Headquarter Location"] for r in records
                           if r.get("Headquarter Location")})

    normalizer = Normalizer()
    industry_map = normalizer.normalize_industries(raw_industries)
    location_map = normalizer.normalize_locations(raw_locations)

    # Print normalization preview.
    print(f"\n[normalize] {len(industry_map)} raw industries → "
          f"{len(set(v for v in industry_map.values() if v))} canonical", flush=True)
    print(f"[normalize] {len(location_map)} raw locations → "
          f"{len(set(str(v) for v in location_map.values()))} canonical", flush=True)

    # Show some interesting merges.
    industry_groups: dict[str, list[str]] = {}
    for raw, canonical in industry_map.items():
        if canonical:
            industry_groups.setdefault(canonical, []).append(raw)
    merged_industries = {k: v for k, v in industry_groups.items() if len(v) > 1}
    if merged_industries:
        print(f"\n[normalize] Industry merges ({len(merged_industries)}):", flush=True)
        for canonical, raws in sorted(merged_industries.items())[:10]:
            print(f"  {raws} → '{canonical}'", flush=True)

    loc_groups: dict[str, list[str]] = {}
    for raw, parsed in location_map.items():
        key = f"{parsed.get('city')}|{parsed.get('state')}|{parsed.get('country')}"
        loc_groups.setdefault(key, []).append(raw)
    merged_locs = {k: v for k, v in loc_groups.items() if len(v) > 1}
    if merged_locs:
        print(f"\n[normalize] Location merges ({len(merged_locs)}):", flush=True)
        for key, raws in sorted(merged_locs.items())[:15]:
            parts = key.split("|")
            print(f"  {raws} → city={parts[0]}, state={parts[1]}, country={parts[2]}",
                  flush=True)

    if dry_run:
        print(f"\n[dry-run] Would insert up to {len(records)} companies. "
              "Pass --execute to actually write to the DB.", flush=True)
        return {"total": len(records), "inserted": 0, "skipped_dup": 0, "skipped_error": 0}

    # 2. Insert.
    now = datetime.now(timezone.utc)
    stats = {"total": len(records), "inserted": 0, "skipped_dup": 0, "skipped_error": 0}

    with conn.cursor() as cur:
        for i, rec in enumerate(records):
            name = (rec.get("Company Name") or "").strip()
            if not name:
                stats["skipped_error"] += 1
                continue

            source = "linkedin"
            platform = (rec.get("Profile URL") or "").strip()
            if not platform:
                stats["skipped_error"] += 1
                continue

            # De-dup: same source + platform = skip.
            if company_exists(cur, source, platform):
                stats["skipped_dup"] += 1
                continue

            # Resolve normalized industry.
            raw_ind = rec.get("Industry")
            canonical_ind = industry_map.get(raw_ind) if raw_ind else None
            industry_id = get_or_create_industry(cur, canonical_ind)

            # Resolve normalized location.
            raw_loc = rec.get("Headquarter Location")
            loc_parsed = location_map.get(raw_loc, {}) if raw_loc else {}
            hq_location_id = get_or_create_location(
                cur, loc_parsed.get("city"), loc_parsed.get("state"),
                loc_parsed.get("country"),
            )

            insert_company(
                cur,
                name=name,
                overview=rec.get("Overview"),
                industry_id=industry_id,
                hq_location_id=hq_location_id,
                logo=rec.get("Logo URL"),
                website=None,
                platform=platform,
                source=source,
                now=now,
            )
            stats["inserted"] += 1

            if (i + 1) % 100 == 0:
                print(f"  [{i + 1}/{len(records)}] inserted so far: {stats['inserted']}",
                      flush=True)

    conn.commit()
    return stats


# ── CLI ───────────────────────────────────────────────────────────

def parse_args(argv=None):
    p = argparse.ArgumentParser(description="Seed PostgreSQL from companies.xlsx")
    p.add_argument("--file", type=str,
                   default=str(PROJECT_ROOT / "companies.xlsx"),
                   help="Path to the Excel file.")
    p.add_argument("--execute", action="store_true",
                   help="Actually insert into the DB (default is dry-run).")
    p.add_argument("--create-tables", action="store_true",
                   help="Run schema.sql to create tables before seeding.")
    p.add_argument("--normalize-only", action="store_true",
                   help="Run LLM normalization, save cache, then exit (no DB needed).")
    p.add_argument("--db-url", type=str, default=None,
                   help="PostgreSQL URL (overrides DATABASE_URL env var).")
    return p.parse_args(argv)


def main(argv=None) -> int:
    import os
    args = parse_args(argv)

    xlsx_path = Path(args.file)
    if not xlsx_path.exists():
        print(f"[error] Excel file not found: {xlsx_path}", file=sys.stderr)
        return 1

    records = read_excel(xlsx_path)
    print(f"[seed] Read {len(records)} rows from {xlsx_path.name}.", flush=True)

    if args.normalize_only:
        raw_industries = sorted({r["Industry"] for r in records if r.get("Industry")})
        raw_locations = sorted({r["Headquarter Location"] for r in records
                               if r.get("Headquarter Location")})
        normalizer = Normalizer()
        ind_map = normalizer.normalize_industries(raw_industries)
        loc_map = normalizer.normalize_locations(raw_locations)
        print(f"\n[done] Normalized {len(ind_map)} industries, {len(loc_map)} locations.",
              flush=True)
        print(f"  Cache saved to .normalization_cache/", flush=True)
        return 0

    db_url = args.db_url or os.getenv("DATABASE_URL")
    if not db_url:
        print("[error] No DATABASE_URL set. Pass --db-url or set it in .env.",
              file=sys.stderr)
        return 1

    conn = psycopg2.connect(db_url)
    try:
        if args.create_tables:
            create_tables(conn)

        stats = seed(records, conn, dry_run=not args.execute)
        print(f"\n[seed] Results: {stats}", flush=True)
    finally:
        conn.close()

    return 0


if __name__ == "__main__":
    sys.exit(main())
